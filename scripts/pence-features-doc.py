"""Write the downloadable Pence feature list. Re-run after any feature change."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / "docs" / "Pence-features.docx"

SECTIONS = [
    (
        "Account",
        [
            "Email sign-in. A new person enters their name, monthly take-home pay, town or city, and payday.",
            "Settings can change those details, including whether spend reminders are wanted.",
            "Settings can save this phone’s location, and other places typed by hand, such as a partner in Essex. Plan and the shop comparison can use any of those places. Location is not required.",
        ],
    ),
    (
        "Adding a spend",
        [
            "Scan a receipt, pick photos from the gallery, or pick PDFs from Files. Up to 8 files. A PDF is read page by page, the first 4 pages.",
            "A receipt, photo or PDF, is read for the shop, date, total, and each line item. You can correct any of it before it is saved. Several receipts are checked one after another.",
            "Pence keeps what it reads, the shop, amount, date and items, and never stores the photo or PDF. While you check a scan, the picture stays on that screen only. Photos saved before this change were cleared.",
            "A photo that is not a receipt or a list of payments is rejected.",
            "A photo or PDF of a bank or card list is accepted. Money in is ignored. Each money-out line can be kept or dropped, with a category.",
            "A wallet or bank alert can be pasted as text. Pence asks whether there is a receipt. If not, it files the spend from the alert.",
            "If a new receipt or statement line matches a payment already saved — same amount, same shop, and the same date — Pence asks whether to keep the saved one, replace it, or keep both. The same amount on a different date is not flagged.",
            "A spend can also be typed in by hand.",
            "When a payment comes in with no receipt, What did you buy? is there to add items, if you want to.",
            "Categories for a logged spend: Groceries, Household, Fashion, Entertainment, Dining, Transport, Health, Bills, Subscriptions, Personal care, Travel, Gifts, and Other.",
        ],
    ),
    (
        "Home",
        [
            "Home shows how much of the current pay is already spent. If payday is the 25th, that pay starts on the 25th of the previous month and runs until the day before the next 25th. With no payday set, it stays a calendar month.",
            "How the month is pacing against last month.",
            "Coming out soon lists regular payments and free trials due in the next 7 days. With none due, a short line opens the full list of regular payments.",
            "Coming up shows the next two buys Pence expects from your own habits, such as a Saturday food shop or a Friday taxi, with the likely day and the usual spend. A buy only shows when there is an offer for it.",
            "Picked for you shows three deals for those habits, each with the reason it is shown. When there are none yet, it shows three posts trending across the UK instead.",
            "Activity on Home lists the latest spends. All Activity opens the full list.",
        ],
    ),
    (
        "Activity",
        [
            "Every spend, grouped by day. A spend can be opened, edited, or deleted.",
            "Editing shows the description as it came in, such as BLD*MYPERFECTCV.CO. A new shop name or category also changes every other payment with the same description, and future ones come in named and filed the same way. Untick it to change only that one.",
            "What did you buy? lets you add items to any payment, with or without prices, such as nappies and bread on a £20.50 ASDA payment with no receipt. Pence learns your habits from them as it does from a receipt. Items without prices share the rest of the total.",
            "A spend read from a receipt says so. The photo itself is not kept.",
        ],
    ),
    (
        "Insights",
        [
            "Charts for where the month went, spending by day, and the past three months side by side.",
            "Each category as a share of take-home, and whether it is up or down.",
            "What to trim, from your own numbers, plus a closer written read when asked.",
            "Another shop lists other big chains actually near this phone, with the distance. A price or an online offer is shown only when a page states it.",
        ],
    ),
    (
        "Offers for you",
        [
            "Pence learns what you buy from every receipt line, every item you add, and every payment: the food shop, coffee, takeaways, clothes, shoes and trainers, perfume, beauty, toiletries, the pharmacy, the home, cleaning, the kids, pets, games, tickets, streaming, CV and job services, other subscriptions, the gym, taxis, trains, fuel, flights, hotels, holidays, tech, gifts, and bills. Perfume bought at Boots counts as perfume, not as the pharmacy. A subscription Pence does not know is not assumed to be streaming.",
            "Pence finds out what each shop really is. It knows the big chains, such as Cake Box for cakes and desserts. Any other payee is looked up on the web, near your town: its real name, what it sells, whether it is a chain, a local business or an online service, where it is, and its website. PIZZA & CO becomes Pizza & Co, a pizza takeaway in your area; BLD*MYPERFECTCV.CO becomes MyPerfectCV, a CV builder. Pence keeps what it learns, and What Pence knows lists your shops with it. Until a shop is looked up, the name is read for a clue, even run together, such as CAKEBOX.",
            "Deals and cheaper places for a shop stay on what it sells: cakes for Cake Box, never burgers. For a local shop, the wider search looks for offers at that shop itself, and nearby places for the same thing.",
            "Before an offer is shown, Pence reads its small print. When the title hints at limits, carries a code, or is an in-store deal, it also reads the deal page. An offer whose terms list places that are not near you is left out, such as a code for a few cities when you live in Essex. An offer only for new customers is left out at a shop you already use. Otherwise the card says what Pence read: Works in Birmingham, the deal page says; Only in some areas, check it works where you live; New customers only; Only some accounts get this. Offers that plainly work for you come first.",
            "When the deal page prints an end date, the card shows it, such as Ends Tue 20 Oct. Deals past their end date are left out.",
            "A deal you can only use in person in a town far from yours, such as a restaurant voucher in Scarborough for someone in Birmingham, is left out. Deals nearby and nationwide stay. Travel deals are kept, since the town is where you are going.",
            "For each habit and each regular shop it works out how often you buy, which weekday, whether it is just after payday, the usual time of day when a receipt shows it, and the usual spend.",
            "From that it expects the next buy, such as a food shop at Tesco on Saturday. Coming up lists those buys with deals for each one. A buy with no offer is left out.",
            "Each habit lists offers at the shops you already use first, marked Your shop. Under them, Could cost less elsewhere shows offers for the same kind of thing at other shops or online, such as Aldi or Lidl when you shop at Tesco, or Vue when you go to Cineworld. Each says which shop you usually use.",
            "For a paid service, such as a CV builder, the wider search also looks for cheaper or free alternatives from public pages.",
            "Deals are grouped by habit. Each one says why it is shown, for example: You buy perfume about every 6 weeks. Deals near your town come first, then the shops you already use, then other shops for the same thing. About thirty at a time, then More.",
            "When a deal page states a percentage off, Pence shows what that is on your usual spend. It never guesses a saving.",
            "A product bought more than once, such as the same perfume, gets a search of its own.",
            "What Pence knows lists every habit. Not me hides one. Add something you buy tells Pence about a habit it has not seen yet, such as holidays. Hidden habits can be shown again.",
            "Deal board results load first. A wider search of public pages follows once a day and is saved for the rest of the day.",
            "Cards and apps for your shops lists the scheme for each shop you use, such as Tesco Clubcard, Nectar, Lidl Plus, Aldi Super 6, or Boots Advantage Card, and a note to check the ride price before you confirm.",
            "Hot deals, opened from Home when there are no personal picks, is the UK trending list, fifteen at a time, up to one hundred when that many exist.",
        ],
    ),
    (
        "Watchlist",
        [
            "For something you want, or somewhere you want to shop, but not at today’s price: a phone, trainers, a hotel, a meal out, ASDA, a gym, any product, service or place.",
            "Added from Plan with Add to Watchlist. It needs the item, or the shop or place, or both. Budget becomes the most you would pay, and the place is used for local offers first.",
            "Pence looks straight away, then every day: HotUKDeals and a wider search of public UK pages. For an item it keeps to the same model and size, leaves out accessories and expired deals, and never takes a monthly contract price for the full price.",
            "For a shop or place on its own, Pence waits for any promotion there. Promotions across the shop come first. A big shop keeps to what you picked in Plan, so ASDA for groceries skips its air fryers.",
            "When an offer turns up, Pence asks whether it is right for you. Taking it ends the watch and keeps the offer under Taken. Still too much keeps watching an item for a lower price than the one turned down; Not this one keeps watching a shop for the next promotion.",
            "Offers seen opens every offer found so far, cheapest first, so you can change your mind and take one at today’s price. A price far below the others, such as a case or a monthly price, is marked as unlikely to be the item and never becomes the price to beat.",
            "Change price clears the price turned down, and offers seen before are checked again against the new price.",
            "The Watchlist is in the Offers tab: offers waiting for an answer first, then everything being watched. Each can be checked now, have its price changed, or be stopped.",
            "Home shows a card when Pence finds an offer, and the Offers tab shows a dot. Phone alerts can be turned on under the Watchlist, and the alert names what was found. On iPhone, Pence must be added to the Home Screen first.",
        ],
    ),
    (
        "Regular payments",
        [
            "A list of subscriptions, bills, memberships and free trials, with when each one is next due and roughly what they cost in a month.",
            "Every payment filed as a subscription is on the list, even one paid once. Other payments join it when they come back on a steady rhythm for about the same amount, such as rent, a gym, or a phone contract.",
            "Payments with the same statement description count as one, so a renamed payee stays one entry.",
            "Pence works out the next date from the last payment. When it has only one payment, it guesses every month and says so. Change fixes the name, amount, how often, and the next date.",
            "Add a payment or free trial covers what the bank does not show yet, such as a trial that starts charging. Once means a free trial or a single payment; it drops off a few days after its date.",
            "A phone alert comes 7 days, 3 days, and 1 day before each payment, naming it and the amount. Each reminder can be turned off. Phone alerts are turned on from the same page.",
            "Send a test shows an alert on this phone straight away, to check alerts arrive. The panel also says when the morning check last ran. If the morning check misses a day, the next visit to Pence runs it.",
            "I’ve cancelled it moves a payment to Cancelled. If it is paid again after that, Pence shows it at the top under Check this.",
            "Not seen lately lists payments that have not come out for longer than usual, to mark as cancelled or still paying.",
            "How to cancel opens a web search for that service.",
            "Opened from Home, from Settings, or from the side menu on a computer.",
        ],
    ),
    (
        "Codes and links",
        [
            "A discount code is shown in full with a Copy button, on every list of deals.",
            "The link goes to the shop where the deal is used. A deal board, Reddit, or a forum is only linked when no shop page is known.",
            "Expired HotUKDeals posts are left out.",
        ],
    ),
    (
        "Plan",
        [
            "The What list on Plan stays on the lower part of the screen. Tap the dim area above it, or the phone’s back button, to close it.",
            "What it is, grouped as Food, Home, Getting around, Shopping, You, and Life. That includes rent, gas and electricity, water, broadband, a mobile contract, TV, council tax, insurance, repairs, furniture, the garden, fuel, parking, a car, a bike, driving lessons, a phone, fashion, gifts, health, personal care, a gym, sport such as a pitch or court, entertainment, hobbies, kids, childcare, pets, education, subscriptions, a wedding or party, moving house, charity, and something else.",
            "A UK plan stays on what was asked for. Deals in the named town come first, then nationwide deals for the same thing. More continues until that list runs out.",
            "A wider search of public pages follows. It is slow the first time. The same plan on the same day is saved and returns faster. A code or a price is shown only when that page prints it.",
            "Saved place comes first. It offers your town, your saved places, Use my location now, and Type a place. Only Type a place shows Country and Town or city.",
            "Shop or place covers shops, restaurants, gyms, pharmacies and anywhere else.",
            "Two buttons: Find offers, and Add to Watchlist.",
        ],
    ),
    (
        "Deal boards",
        [
            "Offers come from anywhere public and legal in the UK: shop sites, Google, Reddit, forums, HotUKDeals, and other UK deal pages. Slickdeals is not used.",
            "LatestDeals, MoneySavingExpert, Groupon, Wowcher, and VoucherCodes are not connected. A person can open them in a browser, but those sites block the app from reading the deals.",
            "The wider offer search can still land on a deal site, Reddit, or a forum when that search actually opens the page.",
        ],
    ),
    (
        "Bank",
        [
            "Not a live bank connection. Google Pay, Apple Pay, PayPal, and Amex are not read automatically.",
            "You can paste or upload a CSV or a copied statement. A photo or PDF of a statement can also be added from Add. Duplicate lines are skipped.",
            "A live bank connection would be a separate Open Banking approval. That is not switched on.",
        ],
    ),
    (
        "Alerts",
        [
            "A pasted wallet alert waits until it is filed or dismissed. That inbox is not in the bottom bar.",
        ],
    ),
    (
        "How you move around",
        [
            "The bottom bar is Home, Plan, Add, Insights, and Offers. The Watchlist is in Offers.",
            "Activity on Home is the latest spends, with All Activity opening the full list. Settings is on Home. Bank import is in Settings.",
            "Regular payments opens from Coming out soon on Home, from Settings, and from the side menu on a computer.",
        ],
    ),
]


def main() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Pence")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x1C, 0x3A, 0x2E)

    sub = doc.add_paragraph()
    sub_run = sub.add_run("Every feature in the current version")
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = RGBColor(0x4A, 0x55, 0x4E)

    dated = doc.add_paragraph()
    dated_run = dated.add_run("11 October 2026")
    dated_run.italic = True
    dated_run.font.size = Pt(11)
    dated_run.font.color.rgb = RGBColor(0x4A, 0x55, 0x4E)

    intro = doc.add_paragraph()
    intro_run = intro.add_run(
        "Pence is a personal spending app in pounds. You sign in and see only your own spends. "
        "It shows where the money went, and it shows an offer only when a public page or a deal board states it."
    )
    intro_run.font.size = Pt(12)

    for heading, lines in SECTIONS:
        h = doc.add_heading(heading, level=1)
        for paragraph in h.runs:
            paragraph.font.color.rgb = RGBColor(0x1C, 0x3A, 0x2E)
        for line in lines:
            bullet = doc.add_paragraph(line, style="List Bullet")
            for paragraph in bullet.runs:
                paragraph.font.size = Pt(12)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
